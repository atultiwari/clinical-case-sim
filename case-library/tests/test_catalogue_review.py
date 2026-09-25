"""The catalogue review pack: build, fill in as Atul would, apply (PLAN L0.5)."""

import csv
import shutil
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts import catalogue as cat
from scripts import catalogue_review as review
from scripts import catalogue_review_apply as apply

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "catalogue"

AUDIT = {
    "range_flags.csv": [
        {
            "component_id": "CMP.HB",
            "sex": "M",
            "severity": "check",
            "issue": "Dacie and Lewis gives 130-170",
            "suggestion": "high 170",
        }
    ],
    "template_flags.csv": [
        {
            "item_id": "EX.ORAL.GUMS",
            "severity": "error",
            "issue": "hint",
            "suggested_text": "Gums healthy.",
        }
    ],
    "cghs_match.csv": [
        {
            "test_id": "LAB.HAEM.CBC",
            "cghs_codes": "LB001",
            "cghs_names": "CBC",
            "cghs_nabh_inr": "175",
            "match": "exact",
            "note": "",
        },
        {
            "test_id": "LAB.TOX.BLOOD_LEAD",
            "cghs_codes": "",
            "cghs_names": "",
            "cghs_nabh_inr": "",
            "match": "none",
            "note": "",
        },
    ],
    "dx_codes.csv": [
        {
            "id": "DX.LEAD_POISONING",
            "icd10_current": "T56.0",
            "icd10_proposed": "T56.0",
            "icd11_current": "",
            "icd11_proposed": "NE61",
            "status": "verified",
            "note": "",
        }
    ],
    "questions.csv": [
        {
            "question_id": "Q1",
            "question": "Blood group default?",
            "options": "O+|B+",
            "recommendation": "B+",
        }
    ],
}


@pytest.fixture
def workspace(tmp_path: Path) -> tuple[Path, Path]:
    """A writable catalogue and an audit folder."""
    catalogue_dir = tmp_path / "catalogue"
    shutil.copytree(FIXTURE, catalogue_dir)
    audit = tmp_path / "audit"
    audit.mkdir()
    for name, rows in AUDIT.items():
        with (audit / name).open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    return catalogue_dir, audit


def _build(workspace: tuple[Path, Path], tmp_path: Path) -> Path:
    catalogue_dir, audit = workspace
    out = tmp_path / "pack.xlsx"
    review.build_pack(cat.read_catalogue(catalogue_dir), audit, out)
    return out


Filler = Callable[[str, dict[str, str]], dict[str, str]]


def _fill(pack: Path, filler: Filler) -> None:
    """Set cells as Atul would: filler(sheet title, row) returns {column: value}."""
    workbook = load_workbook(pack)
    for sheet in review.SHEETS:
        ws = workbook[sheet.title]
        header = [c.value for c in ws[1]]
        for row_cells in ws.iter_rows(min_row=2):
            row = {
                str(h): review.cell_text(c.value) for h, c in zip(header, row_cells, strict=True)
            }
            for column, value in filler(sheet.title, row).items():
                row_cells[header.index(column)].value = value
    workbook.save(pack)


def _approve_all(title: str, _: dict[str, str]) -> dict[str, str]:
    return {"decision": "Keep B+" if title == "Questions" else "approve"}


def _csv(directory: Path, name: str) -> list[dict[str, str]]:
    with (directory / name).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def test_build_writes_every_sheet_with_flagged_rows_first(
    workspace: tuple[Path, Path], tmp_path: Path
) -> None:
    pack = _build(workspace, tmp_path)
    workbook = load_workbook(pack)
    assert workbook.sheetnames == ["Read me", *(s.title for s in review.SHEETS)]
    ranges = workbook["Ranges"]
    header = [c.value for c in ranges[1]]
    first = dict(zip(header, (c.value for c in ranges[2]), strict=True))
    assert (first["component_id"], first["sex"], first["flag"]) == ("CMP.HB", "M", "check")
    assert ranges.max_row == 1 + 7
    texts = workbook["Normal texts"]
    assert texts.cell(2, 1).value == "EX.ORAL.GUMS"
    assert texts.max_row == 1 + 4 + 2  # four templates, two qualitative normal texts
    prices = {
        r[0]: r for r in workbook["Prices and turnaround"].iter_rows(min_row=2, values_only=True)
    }
    assert prices["LAB.HAEM.CBC"][9:11] == ("175", review.CGHS_SOURCE)
    assert prices["LAB.TOX.BLOOD_LEAD"][9:11] == ("1500", "estimate")


def test_an_undecided_pack_is_refused_and_nothing_is_written(
    workspace: tuple[Path, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    pack = _build(workspace, tmp_path)
    before = (workspace[0] / "tests.csv").read_text(encoding="utf-8")
    _fill(
        pack, lambda title, row: {} if row.get("ref") == "CMP.RBC#1" else _approve_all(title, row)
    )
    assert apply.run_apply(pack, workspace[0], dry_run=False) == 1
    assert "Ranges CMP.RBC#1: no decision" in capsys.readouterr().err
    assert (workspace[0] / "tests.csv").read_text(encoding="utf-8") == before


@pytest.mark.parametrize(
    ("sheet", "decision", "expected"),
    [
        ("Ranges", "maybe", "is not one of"),
        ("Ranges", "edit", "no New column"),
        ("Normal texts", "suggestion", "no suggested text"),
    ],
)
def test_invalid_decisions_are_refused(
    workspace: tuple[Path, Path], tmp_path: Path, sheet: str, decision: str, expected: str
) -> None:
    pack = _build(workspace, tmp_path)
    target = {"Ranges": "CMP.RBC#1", "Normal texts": "HX.EXPOSURE.OCCUPATION"}[sheet]

    def filler(title: str, row: dict[str, str]) -> dict[str, str]:
        key = row.get("ref") or row.get("item_id")
        return (
            {"decision": decision} if title == sheet and key == target else _approve_all(title, row)
        )

    _fill(pack, filler)
    with pytest.raises(apply.ReviewError, match=expected):
        apply.apply_decisions(cat.read_catalogue(workspace[0]), apply.read_decisions(pack))


def test_approving_everything_applies_the_proposals(
    workspace: tuple[Path, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    pack = _build(workspace, tmp_path)
    _fill(pack, _approve_all)
    assert apply.run_apply(pack, workspace[0], dry_run=False) == 0
    out = capsys.readouterr().out
    assert "Question Q1: Blood group default?" in out
    assert "Answer: Keep B+" in out
    tests = {r["id"]: r for r in _csv(workspace[0], "tests.csv")}
    assert (tests["LAB.HAEM.CBC"]["price_inr"], tests["LAB.HAEM.CBC"]["price_source"]) == (
        "175",
        review.CGHS_SOURCE,
    )
    assert tests["LAB.TOX.BLOOD_LEAD"]["price_source"] == "estimate"
    templates = _csv(workspace[0], "normal_templates.csv")
    assert {t["review_status"] for t in templates} == {"approved"}
    assert next(t for t in templates if t["item_id"] == "EX.ORAL.GUMS")["template"] == (
        "Gum margins normal; no pigmentation."
    )
    dx = _csv(workspace[0], "diagnoses.csv")[0]
    assert (dx["icd11"], dx["codes_verified"]) == ("NE61", "true")
    assert cat.check(cat.read_catalogue(workspace[0])) == []


def test_edits_and_suggestions_are_applied(workspace: tuple[Path, Path], tmp_path: Path) -> None:
    pack = _build(workspace, tmp_path)
    edits = {
        "CMP.HB#2": {"decision": "edit", "new_high": "170"},
        "CMP.PB_BLOOD#1": {"decision": "edit", "new_high": "3.5", "new_display": "<3.5"},
        "CMP.MCH#1": {"decision": "edit", "new_decimals": "0"},
        "EX.ORAL.GUMS": {"decision": "suggestion"},
        "CMP.FILM_REPORT": {"decision": "edit", "new_text": "Normal film."},
        "LAB.TOX.BLOOD_LEAD": {
            "decision": "edit",
            "new_price_inr": "900",
            "new_tat_minutes": "2880",
        },
        "DX.LEAD_POISONING": {"decision": "unverified"},
    }

    def filler(title: str, row: dict[str, str]) -> dict[str, str]:
        key = row.get("ref") or row.get("item_id") or row.get("test_id") or row.get("dx_id")
        return edits.get(key or "", _approve_all(title, row))

    _fill(pack, filler)
    files, outcome = apply.apply_decisions(
        cat.read_catalogue(workspace[0]), apply.read_decisions(pack)
    )
    ranges = files["component_ranges.csv"]
    assert ranges[1]["high"] == "170"
    lead = next(r for r in ranges if r["component_id"] == "CMP.PB_BLOOD")
    assert (lead["high"], lead["display"]) == ("3.5", "<3.5")
    components = {r["id"]: r for r in files["components.csv"]}
    assert components["CMP.MCH"]["decimals"] == "0"
    assert components["CMP.FILM_REPORT"]["normal_text"] == "Normal film."
    gums = next(t for t in files["normal_templates.csv"] if t["item_id"] == "EX.ORAL.GUMS")
    assert (gums["template"], gums["review_status"]) == ("Gums healthy.", "approved")
    lead_test = next(t for t in files["tests.csv"] if t["id"] == "LAB.TOX.BLOOD_LEAD")
    assert (lead_test["price_inr"], lead_test["price_source"], lead_test["tat_minutes"]) == (
        "900",
        "estimate",
        "2880",
    )
    assert files["diagnoses.csv"][0]["codes_verified"] == "false"
    assert outcome.changed["component_ranges.csv"] == 2


def test_a_dry_run_writes_nothing(
    workspace: tuple[Path, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    pack = _build(workspace, tmp_path)
    _fill(pack, _approve_all)
    before = (workspace[0] / "tests.csv").read_text(encoding="utf-8")
    assert apply.run_apply(pack, workspace[0], dry_run=True) == 0
    assert "Would change: tests.csv 1 rows" in capsys.readouterr().out
    assert (workspace[0] / "tests.csv").read_text(encoding="utf-8") == before


def test_cli_build(workspace: tuple[Path, Path], tmp_path: Path) -> None:
    out = tmp_path / "cli.xlsx"
    args = [
        "build",
        "--version",
        "1",
        "--dir",
        str(workspace[0]),
        "--audit",
        str(workspace[1]),
        "--out",
        str(out),
    ]
    assert review.main(args) == 0
    assert out.exists()
    _fill(out, _approve_all)
    assert review.main(["apply", str(out), "--dir", str(workspace[0]), "--dry-run"]) == 0


def test_a_workbook_missing_a_sheet_is_refused(
    workspace: tuple[Path, Path], tmp_path: Path
) -> None:
    pack = _build(workspace, tmp_path)
    workbook = load_workbook(pack)
    del workbook["Questions"]
    workbook.save(pack)
    with pytest.raises(apply.ReviewError, match="Questions"):
        apply.read_decisions(pack)
