"""The case review pack: write it from plain rows, fill it in as Atul would, read it (SPEC §7.2)."""

import json
import subprocess
import sys
from collections.abc import Callable
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts import review_pack
from scripts import review_pack_data as data
from scripts import review_pack_read as reader
from scripts import review_pack_write as writer

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "review_pack_data.json"
CASE_LIBRARY = Path(__file__).resolve().parents[1]
SHEET_TITLES = [
    "README",
    "Judgement calls",
    "Affected",
    "Normal list",
    "Article facts",
    "Patient's words",
    "Reports and consults",
    "Ground truth",
    "Leak scan",
]
Filler = Callable[[str, dict[str, str]], dict[str, str]]


@pytest.fixture
def pack_data() -> data.PackData:
    return data.load_pack_data(FIXTURE)


@pytest.fixture
def pack(tmp_path: Path, pack_data: data.PackData) -> Path:
    out = tmp_path / "fixture-batch.xlsx"
    writer.write_pack(pack_data, out)
    return out


def _rows(path: Path, title: str) -> list[dict[str, str]]:
    ws = load_workbook(path)[title]
    values = list(ws.iter_rows(values_only=True))
    header = [str(v) for v in values[0]]
    return [
        {h: "" if v is None else str(v) for h, v in zip(header, row, strict=False)}
        for row in values[1:]
    ]


def _fill(path: Path, choose: Filler) -> None:
    """Write Atul's decisions: choose(sheet title, row) -> {column: value}."""
    workbook = load_workbook(path)
    for title in SHEET_TITLES[1:]:
        ws = workbook[title]
        header = [c.value for c in ws[1]]
        for row_index in range(2, ws.max_row + 1):
            row = {
                str(h): "" if c.value is None else str(c.value)
                for h, c in zip(header, ws[row_index], strict=False)
            }
            for column, value in choose(title, row).items():
                ws.cell(row_index, header.index(column) + 1).value = value
    workbook.save(path)


def _default_choice(title: str, row: dict[str, str]) -> dict[str, str]:
    if title in {"Judgement calls", "Affected", "Reports and consults"}:
        return {"decision": "Approve"}
    if title == "Leak scan":
        return {"decision": "Resolved"}
    return {"decision": "OK"}


# --- writing -------------------------------------------------------------------------------


def test_write_pack_has_every_sheet_and_counts(tmp_path: Path, pack_data: data.PackData) -> None:
    out = tmp_path / "pack.xlsx"
    counts = writer.write_pack(pack_data, out)
    assert load_workbook(out).sheetnames == SHEET_TITLES
    assert counts == {
        "Judgement calls": 1,
        "Affected": 2,
        "Normal list": 2,
        "Article facts": 2,
        "Patient's words": 1,
        "Reports and consults": 2,
        "Ground truth": 3,
        "Leak scan": 1,
    }


def test_judgement_calls_are_separate_and_high_priority_first(pack: Path) -> None:
    judgement = _rows(pack, "Judgement calls")
    assert [r["target"] for r in judgement] == ["CMP.ZPP"]
    affected = _rows(pack, "Affected")
    assert [r["priority"] for r in affected] == ["high", "low"]
    assert affected[0]["key"].startswith("synthetic_ledger:")


def test_every_row_has_a_key_and_a_decision_dropdown(pack: Path) -> None:
    workbook = load_workbook(pack)
    for title in SHEET_TITLES[1:]:
        ws = workbook[title]
        header = [c.value for c in ws[1]]
        assert header[0] == "key"
        assert "decision" in header
        assert all(ws.cell(r, 1).value for r in range(2, ws.max_row + 1))
        formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        assert formulas, title
    lists = {t: workbook[t].data_validations.dataValidation[0].formula1 for t in SHEET_TITLES[1:]}
    assert lists["Affected"] == '"Approve,Edit,Reject"'
    assert lists["Normal list"] == '"OK,Should be affected"'
    assert lists["Article facts"] == '"OK,Correction"'
    assert lists["Leak scan"] == '"Resolved"'


def test_edit_columns_are_offered(pack: Path) -> None:
    affected = _rows(pack, "Affected")[0]
    assert {"new_value", "new_text", "note"} <= set(affected)
    assert "new_text" in _rows(pack, "Reports and consults")[0]
    assert "correction" in _rows(pack, "Ground truth")[0]


def test_normal_list_shows_names_not_values(pack: Path) -> None:
    rows = _rows(pack, "Normal list")
    assert [r["item"] for r in rows] == ["Sodium", "Occupation"]
    assert "value" not in rows[0]


def test_readme_names_batch_cases_counts_and_licence(pack: Path) -> None:
    ws = load_workbook(pack)["README"]
    text = "\n".join(
        " | ".join("" if v is None else str(v) for v in row)
        for row in ws.iter_rows(values_only=True)
    )
    assert "fixture-batch" in text
    assert "NID-9001@v1" in text
    assert "de novo (no article)" in text
    assert "n/a" in text  # licence flags do not apply to a de novo case
    assert "affected" in text
    assert "Judgement calls" in text


def test_readme_licence_flags_for_an_article(tmp_path: Path, pack_data: data.PackData) -> None:
    case = data.CaseSummary("PMC1@v1", "draft", "PMC1", "CC BY 4.0", True, False, {"article": 1})
    out = tmp_path / "p.xlsx"
    writer.write_pack(data.PackData("b", (case,), (), (), (), (), (), (), ()), out)
    rows = list(load_workbook(out)["README"].iter_rows(values_only=True))
    case_row = next(r for r in rows if r[0] == "PMC1@v1")
    assert "yes" in case_row
    assert "no" in case_row
    assert all(load_workbook(out)[t].max_row == 1 for t in SHEET_TITLES[1:])


# --- pure helpers --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ({"value": 75}, "75"),
        ({"value": 2.5, "unit": "x10^12/L"}, "2.5 x10^12/L"),
        ({"text": "Straw"}, "Straw"),
        ({"a": 1}, '{"a": 1}'),
        ("plain", "plain"),
        (None, ""),
    ],
)
def test_format_value(value: object, expected: str) -> None:
    assert data.format_value(value) == expected


def test_format_range_and_text() -> None:
    assert data.format_range(Decimal("115"), Decimal("165"), "g/L") == "115-165 g/L"
    assert data.format_range(None, None, None) == ""
    assert data.as_text(Decimal("0.70")) == "0.7"
    assert data.as_text(None) == ""
    assert data.as_text(True) == "yes"
    assert data.as_text(3) == "3"


def test_ground_truth_rows_flatten_rubric_and_conditions() -> None:
    gt = {
        "final_dx": {"id": "DX.X", "text": "The diagnosis"},
        "rubric": {
            "rubric": [{"score": 5, "text": "Best", "if": {"dx_in": ["DX.X"]}}],
            "default_score": 1,
        },
        "must_do": [{"text": "Do it", "if": {"ordered_any": ["LAB.A"]}}, "Plain wording"],
        "must_not_do": [],
    }
    rows = data.ground_truth_rows("C@v1", gt)
    assert [r["key"] for r in rows] == [
        "ground_truth:C@v1/final_dx",
        "ground_truth:C@v1/rubric/1",
        "ground_truth:C@v1/must_do/1",
        "ground_truth:C@v1/must_do/2",
    ]
    assert rows[0]["text"] == "The diagnosis"
    assert rows[0]["condition"] == '{"id": "DX.X"}'
    assert rows[1]["text"] == "Score 5: Best"
    assert rows[1]["condition"] == '{"dx_in": ["DX.X"]}'
    assert rows[3] == {
        "key": "ground_truth:C@v1/must_do/2",
        "case": "C@v1",
        "part": "must_do",
        "text": "Plain wording",
        "condition": "",
    }


# --- reading -------------------------------------------------------------------------------


def test_read_round_trip_writes_decisions(pack: Path, capsys: pytest.CaptureFixture[str]) -> None:
    def choose(title: str, row: dict[str, str]) -> dict[str, str]:
        if row["key"].endswith("0002"):
            return {"decision": "Edit", "new_value": "78 g/L", "note": "Closer to day 2"}
        if row["key"].endswith("0001"):
            return {"decision": "Reject"}
        if title == "Article facts" and row["id"] == "L01":
            return {"decision": "Correction", "correction": "2.45"}
        if title == "Normal list" and row["target"] == "CMP.NA":
            return {"decision": "Should be affected"}
        return _default_choice(title, row)

    _fill(pack, choose)
    assert review_pack.main(["read", str(pack)]) == 0
    out = capsys.readouterr().out
    assert "NID-9001@v1" in out
    assert "approved 3" in out  # the judgement call, the report and the consult note
    assert "edited 1" in out
    assert "rejected 1" in out
    assert "corrections 2" in out  # the article fact and the normal item

    decisions = json.loads(pack.with_name("fixture-batch.decisions.json").read_text("utf-8"))
    assert len(decisions) == 14
    edit = next(d for d in decisions if d["decision"] == "edit")
    assert edit == {
        "key": "synthetic_ledger:00000000-0000-4000-8000-000000000002",
        "sheet": "Affected",
        "table": "synthetic_ledger",
        "id": "00000000-0000-4000-8000-000000000002",
        "case_version_id": "NID-9001@v1",
        "decision": "edit",
        "edited": {"new_value": "78 g/L"},
        "note": "Closer to day 2",
    }
    fact = next(d for d in decisions if d["decision"] == "correction" and d["table"] == "fact")
    assert fact["id"] == "NID-9001@v1/L01"
    assert fact["edited"] == {"correction": "2.45"}
    assert {d["decision"] for d in decisions} >= {"should_be_affected", "resolved", "ok"}
    approve = next(d for d in decisions if d["decision"] == "approve")
    assert approve["edited"] is None
    assert approve["note"] == ""


def test_read_refuses_undecided_and_invalid_rows(
    pack: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    def choose(title: str, row: dict[str, str]) -> dict[str, str]:
        if title == "Leak scan":
            return {}
        if title == "Normal list":
            return {"decision": "Maybe"}
        return _default_choice(title, row)

    _fill(pack, choose)
    assert review_pack.main(["read", str(pack)]) == 1
    err = capsys.readouterr().err
    assert "Leak scan leak_scan:NID-9001@v1/report/RP01/plumbism: no decision" in err
    assert err.count("'Maybe'") == 2
    assert not pack.with_name("fixture-batch.decisions.json").exists()


def test_read_refuses_edit_without_new_value(
    pack: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    def choose(title: str, row: dict[str, str]) -> dict[str, str]:
        if title == "Reports and consults" and row["kind"] == "report":
            return {"decision": "Edit"}
        if title == "Ground truth" and row["part"] == "must_do":
            return {"decision": "Correction"}
        return _default_choice(title, row)

    _fill(pack, choose)
    assert review_pack.main(["read", str(pack)]) == 1
    err = capsys.readouterr().err
    assert "report:NID-9001@v1/RP01: Edit but no new_text" in err
    assert "ground_truth:NID-9001@v1/must_do/1: Correction but no correction" in err
    assert not pack.with_name("fixture-batch.decisions.json").exists()


def test_read_accepts_any_letter_case(pack: Path) -> None:
    _fill(pack, lambda title, row: {"decision": _default_choice(title, row)["decision"].lower()})
    assert review_pack.main(["read", str(pack)]) == 0


def test_read_refuses_a_workbook_without_a_sheet(
    tmp_path: Path, pack: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = load_workbook(pack)
    workbook.remove(workbook["Leak scan"])
    broken = tmp_path / "broken.xlsx"
    workbook.save(broken)
    assert review_pack.main(["read", str(broken)]) == 1
    assert "no 'Leak scan' sheet" in capsys.readouterr().err


def test_read_refuses_a_missing_file(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert review_pack.main(["read", str(tmp_path / "nope.xlsx")]) == 1
    assert "not found" in capsys.readouterr().err


# --- command line --------------------------------------------------------------------------


def test_build_without_database_url_explains(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.delenv("CASE_VAULT_DB_URL_READONLY", raising=False)
    assert review_pack.main(["build", "b1", "--cases", "NID-9001@v1"]) == 1
    assert "CASE_VAULT_DB_URL_READONLY is not set" in capsys.readouterr().err


def test_default_output_path() -> None:
    assert review_pack.default_out("pilot") == CASE_LIBRARY / "review" / "pilot.xlsx"


def test_batch_name_is_checked(capsys: pytest.CaptureFixture[str]) -> None:
    assert review_pack.main(["build", "../escape", "--cases", "X@v1"]) == 1
    assert "batch name" in capsys.readouterr().err


def test_runs_as_a_file() -> None:
    result = subprocess.run(
        [sys.executable, str(CASE_LIBRARY / "scripts" / "review_pack.py"), "--help"],
        capture_output=True,
        text=True,
        check=False,
        cwd=CASE_LIBRARY.parent,
    )
    assert result.returncode == 0, result.stderr
    assert "build" in result.stdout


def test_decisions_sit_beside_the_pack() -> None:
    assert reader.decisions_path(Path("review/pilot.xlsx")) == Path("review/pilot.decisions.json")
