"""The review pack builder on the fixture case in the local database (SPEC §7.2, PLAN L0.6)."""

from pathlib import Path

import psycopg
import pytest
from openpyxl import load_workbook

from scripts import review_pack
from scripts.review_pack_data import ReviewPackError, fetch_pack_data
from scripts.review_pack_write import write_pack
from tests.conftest import LOCAL_DB_URL
from tests.integration.conftest import MINI_GENERATOR, add_affected_cbc, seed_mini_case

pytestmark = pytest.mark.integration


def _curate(db: psycopg.Connection) -> str:
    """The fixture case part-way through curation, with one of everything the pack shows."""
    cv = seed_mini_case(db)
    db.execute("select casevault.compute_derived(%s, %s, 'v0.0')", (cv, MINI_GENERATOR))
    add_affected_cbc(db, cv)
    db.execute(
        "insert into casevault.synthetic_ledger (case_version_id, target, day_bucket, tier, value,"
        " rationale, confidence, priority, judgement_call, generator, skill_version) values"
        " (%(cv)s, 'CMP.PB', 1, 'affected', '{\"value\": 70}', 'No anchor', 0.5, 'high', true,"
        "  %(g)s, 'v0'),"
        " (%(cv)s, 'HX.OCCUPATION', null, 'affected', '{\"text\": \"Shop keeper\"}', 'Job', 0.9,"
        "  'high', false, %(g)s, 'v0')",
        {"cv": cv, "g": MINI_GENERATOR},
    )
    db.execute(
        "update casevault.fact set lay_text = 'I take a herbal tonic.'"
        " where case_version_id = %s and id = 'H02'",
        (cv,),
    )
    db.execute("select casevault.resolve_normals(%s)", (cv,))
    db.execute(
        "insert into casevault.report (case_version_id, id, test_item_id, variant, status,"
        " status_line, findings, report_text, generator, skill_version) values"
        " (%s, 'RP01', 'LAB.FILM', 'original', 'provisional', 'Provisional: not final.',"
        " '{FND.STIPPLING}', 'Stippling; think of plumbism.', 'g', 's')",
        (cv,),
    )
    db.execute(
        "insert into casevault.consult_note (case_version_id, id, specialty, condition, note_text,"
        " origin, rationale, confidence, generator, skill_version) values"
        " (%s, 'CN01', 'REF.TOXICOLOGY', '{\"not\": {\"released_any\": [\"L02\"]}}',"
        " 'Happy to see her.', 'affected', 'r', 0.8, 'g', 's')",
        (cv,),
    )
    return cv


def _count(db: psycopg.Connection, sql: str, cv: str) -> int:
    row = db.execute(sql, (cv,)).fetchone()
    assert row is not None
    return int(row[0])


def test_builder_writes_the_pack_from_the_fixture(db: psycopg.Connection, tmp_path: Path) -> None:
    cv = _curate(db)
    pack = fetch_pack_data(db, [cv], batch="fixture-batch")

    (case,) = pack.cases
    assert case.case_version_id == cv
    assert case.source == "PMC0000001"
    assert case.licence == "CC BY 4.0"
    assert (case.production_ok, case.public_release_ok) == (True, True)
    # 5 written here, plus MCH on days 1 and 2, derived by the generator from affected inputs
    assert case.origin_counts["affected"] == 7
    assert case.origin_counts["article"] == 6
    assert case.origin_counts["derived"] >= 1

    out = tmp_path / "fixture-batch.xlsx"
    counts = write_pack(pack, out)
    normals = _count(
        db,
        "select count(distinct target) from casevault.synthetic_ledger"
        " where case_version_id = %s and tier in ('normal', 'rule')",
        cv,
    )
    assert counts == {
        "Judgement calls": 1,
        "Affected": 6,
        "Normal list": normals,
        "Article facts": 6,
        "Patient's words": 2,  # H01 and H02 are the history facts
        "Reports and consults": 2,
        "Ground truth": 2,  # the final diagnosis and one must-do
        "Leak scan": 1,
    }
    assert normals > 0

    workbook = load_workbook(out)
    judgement = list(workbook["Judgement calls"].iter_rows(min_row=2, values_only=True))
    assert judgement[0][2] == "CMP.PB"
    assert judgement[0][6] == "70"
    assert judgement[0][7] == "0-5 ug/dL"
    affected = list(workbook["Affected"].iter_rows(min_row=2, values_only=True))
    assert affected[0][2] == "HX.OCCUPATION"  # high priority first
    assert affected[0][6] == "Shop keeper"
    leak = list(workbook["Leak scan"].iter_rows(min_row=2, values_only=True))
    assert leak[0][0] == f"leak_scan:{cv}/report/RP01/plumbism"
    words = [r[0] for r in workbook["Patient's words"].iter_rows(min_row=2, values_only=True)]
    assert words == [f"fact:{cv}/H01", f"fact:{cv}/H02"]


def test_builder_refuses_unknown_case_versions(db: psycopg.Connection) -> None:
    cv = seed_mini_case(db)
    with pytest.raises(ReviewPackError, match="NID-0000@v9"):
        fetch_pack_data(db, [cv, "NID-0000@v9"])


def test_command_line_build_reports_unknown_cases(
    local_db: psycopg.Connection,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setenv("CASE_VAULT_DB_URL_READONLY", LOCAL_DB_URL)
    out = tmp_path / "b.xlsx"
    assert review_pack.main(["build", "b", "--cases", "NID-0000@v9", "--out", str(out)]) == 1
    assert "No such case version in the Case Vault: NID-0000@v9" in capsys.readouterr().err
    assert not out.exists()
