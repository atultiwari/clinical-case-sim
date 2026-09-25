"""Build the catalogue review pack, an Excel workbook for Atul (PLAN L0.5).

    uv run python -m scripts.catalogue_review build --version 1
    uv run python -m scripts.catalogue_review apply review/catalogue-v1/catalogue-review-v1.xlsx

`build` reads catalogue/*.csv and the audit files in review/catalogue-v<n>/audit/
(CGHS matches, range, template and code flags, open questions) and writes one
sheet per review area. Flagged rows come first. Atul fills the yellow Decision
column (and the New columns for an edit); `apply` (scripts/catalogue_review_apply.py)
writes the decisions back into the CSV files. Review packs stay out of Git.
"""

import argparse
import csv
import sys
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from scripts.catalogue import CASE_LIBRARY, DEFAULT_DIR, Catalogue, read_catalogue
from scripts.workbook import FILLS, cell_text, write_table

__all__ = ["FILLS", "SHEETS", "build_pack", "cell_text", "main", "range_refs"]

REVIEW_DIR = CASE_LIBRARY / "review"
CGHS_SOURCE = "CGHS 2025, Tier I NABH (OM 03.10.2025)"
SEVERITY_ORDER = {"error": 0, "check": 1, "minor": 2, "": 3}


@dataclass(frozen=True)
class Sheet:
    """One review sheet: shown columns, then Decision and the New columns an edit fills."""

    title: str
    key: str
    columns: tuple[str, ...]
    decisions: tuple[str, ...]  # empty = free text (Questions)
    edits: tuple[str, ...]


SHEETS: tuple[Sheet, ...] = (
    Sheet(
        "Ranges",
        "ref",
        (
            "ref",
            "component_id",
            "name",
            "unit_si",
            "decimals",
            "unit_conv",
            "conv_factor",
            "sex",
            "age_min",
            "age_max",
            "low",
            "high",
            "text",
            "display",
            "source",
            "flag",
            "issue",
            "suggestion",
        ),
        ("approve", "edit"),
        (
            "new_low",
            "new_high",
            "new_display",
            "new_text",
            "new_source",
            "new_unit_si",
            "new_decimals",
            "new_conv_factor",
        ),
    ),
    Sheet(
        "Normal texts",
        "item_id",
        ("item_id", "kind", "name", "text", "flag", "issue", "suggestion"),
        ("approve", "suggestion", "edit"),
        ("new_text",),
    ),
    Sheet(
        "Prices and turnaround",
        "test_id",
        (
            "test_id",
            "name",
            "route",
            "estimate_inr",
            "cghs_codes",
            "cghs_names",
            "cghs_nabh_inr",
            "match",
            "match_note",
            "proposed_inr",
            "proposed_source",
            "tat_minutes",
            "flag",
        ),
        ("approve", "edit"),
        ("new_price_inr", "new_price_source", "new_tat_minutes"),
    ),
    Sheet(
        "Diagnosis codes",
        "dx_id",
        (
            "dx_id",
            "name",
            "icd10",
            "icd11",
            "proposed_icd10",
            "proposed_icd11",
            "status",
            "audit_note",
            "flag",
        ),
        ("approve", "edit", "unverified"),
        ("new_icd10", "new_icd11"),
    ),
    Sheet(
        "Questions", "question_id", ("question_id", "question", "options", "recommendation"), (), ()
    ),
)
SHEET_BY_TITLE: Mapping[str, Sheet] = {s.title: s for s in SHEETS}

README_LINES = (
    "Catalogue review pack (Case Library, PLAN L0.5). Educational and research use only.",
    "",
    "Every row needs a Decision (the yellow column). Flagged rows come first: red = error, "
    "amber = check. 'issue' and 'suggestion' are Claude's audit notes, not decisions.",
    "",
    "Ranges: approve = keep as shown; edit = fill only the New columns you change.",
    "Normal texts: approve = keep; suggestion = use the suggested text; edit = write new_text. "
    "Approved and edited history, exam and referral templates load as approved.",
    "Prices and turnaround: approve = take proposed_inr and proposed_source (the CGHS NABH "
    "Tier I rate where one matches, otherwise the estimate) and tat_minutes; edit = fill the "
    "New columns you change.",
    "Diagnosis codes: approve = take the proposed codes as verified; edit = new codes, "
    "verified; unverified = keep the proposed codes, still marked unverified.",
    "Questions: write your answer in Decision.",
    "",
    "Fill-down is fine for a block of rows you accept. Return the file as it is; Claude runs "
    "`python -m scripts.catalogue_review apply` on it.",
)


def read_optional_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8", newline="") as handle:
        return [{k: (v or "").strip() for k, v in row.items()} for row in csv.DictReader(handle)]


def range_refs(catalogue: Catalogue) -> list[tuple[str, Mapping[str, str]]]:
    """(ref, range row) with ref = '<component>#<n>', n counting that component's rows from 1."""
    counts: dict[str, int] = {}
    refs = []
    for row in catalogue.rows("component_ranges.csv"):
        counts[row["component_id"]] = counts.get(row["component_id"], 0) + 1
        refs.append((f"{row['component_id']}#{counts[row['component_id']]}", row))
    return refs


def _range_rows(catalogue: Catalogue, audit: Path) -> list[dict[str, str]]:
    components = {r["id"]: r for r in catalogue.rows("components.csv")}
    flags: dict[tuple[str, str], dict[str, str]] = {}
    for flag in read_optional_csv(audit / "range_flags.csv"):
        flags[(flag["component_id"], flag["sex"])] = flag
    rows = []
    for ref, rng in range_refs(catalogue):
        comp = components[rng["component_id"]]
        flag = flags.get((rng["component_id"], rng["sex"])) or flags.get(
            (rng["component_id"], ""), {}
        )
        rows.append(
            {
                "ref": ref,
                "component_id": rng["component_id"],
                "name": comp["name"],
                "unit_si": comp["unit_si"],
                "decimals": comp["decimals"],
                "unit_conv": comp["unit_conv"],
                "conv_factor": comp["conv_factor"],
                **{
                    k: rng[k]
                    for k in (
                        "sex",
                        "age_min",
                        "age_max",
                        "low",
                        "high",
                        "text",
                        "display",
                        "source",
                    )
                },
                "flag": flag.get("severity", ""),
                "issue": flag.get("issue", ""),
                "suggestion": flag.get("suggestion", ""),
            }
        )
    return rows


def _text_rows(catalogue: Catalogue, audit: Path) -> list[dict[str, str]]:
    flags = {f["item_id"]: f for f in read_optional_csv(audit / "template_flags.csv")}
    names = {r["id"]: (kind, r["name"]) for kind, r in catalogue.items()}
    rows = [
        {
            "item_id": t["item_id"],
            "kind": names[t["item_id"]][0],
            "name": names[t["item_id"]][1],
            "text": t["template"],
        }
        for t in catalogue.rows("normal_templates.csv")
    ]
    rows += [
        {"item_id": c["id"], "kind": "component", "name": c["name"], "text": c["normal_text"]}
        for c in catalogue.rows("components.csv")
        if c["normal_text"]
    ]
    for row in rows:
        flag = flags.get(row["item_id"], {})
        row.update(
            flag=flag.get("severity", ""),
            issue=flag.get("issue", ""),
            suggestion=flag.get("suggested_text", ""),
        )
    return rows


def _price_rows(catalogue: Catalogue, audit: Path) -> list[dict[str, str]]:
    matches = {m["test_id"]: m for m in read_optional_csv(audit / "cghs_match.csv")}
    rows = []
    for test in catalogue.rows("tests.csv"):
        match = matches.get(test["id"], {})
        has_rate = match.get("match") in {"exact", "close"} and bool(match.get("cghs_nabh_inr"))
        rows.append(
            {
                "test_id": test["id"],
                "name": test["name"],
                "route": test["route"],
                "estimate_inr": test["price_inr"],
                "cghs_codes": match.get("cghs_codes", ""),
                "cghs_names": match.get("cghs_names", ""),
                "cghs_nabh_inr": match.get("cghs_nabh_inr", ""),
                "match": match.get("match", ""),
                "match_note": match.get("note", ""),
                "proposed_inr": match["cghs_nabh_inr"] if has_rate else test["price_inr"],
                "proposed_source": CGHS_SOURCE if has_rate else "estimate",
                "tat_minutes": test["tat_minutes"],
                "flag": "" if has_rate else "check",
            }
        )
    return rows


def _dx_rows(catalogue: Catalogue, audit: Path) -> list[dict[str, str]]:
    codes = {c["id"]: c for c in read_optional_csv(audit / "dx_codes.csv")}
    rows = []
    for dx in catalogue.rows("diagnoses.csv"):
        code = codes.get(dx["id"], {})
        status = code.get("status", "unverified")
        rows.append(
            {
                "dx_id": dx["id"],
                "name": dx["name"],
                "icd10": dx["icd10"],
                "icd11": dx["icd11"],
                "proposed_icd10": code.get("icd10_proposed") or dx["icd10"],
                "proposed_icd11": code.get("icd11_proposed", dx["icd11"]),
                "status": status,
                "audit_note": code.get("note", ""),
                "flag": {"changed": "error", "unverified": "check"}.get(status, ""),
            }
        )
    return rows


def _question_rows(_: Catalogue, audit: Path) -> list[dict[str, str]]:
    return read_optional_csv(audit / "questions.csv")


ROW_BUILDERS: Mapping[str, Callable[[Catalogue, Path], list[dict[str, str]]]] = {
    "Ranges": _range_rows,
    "Normal texts": _text_rows,
    "Prices and turnaround": _price_rows,
    "Diagnosis codes": _dx_rows,
    "Questions": _question_rows,
}


def _write_sheet(ws: Worksheet, sheet: Sheet, rows: list[dict[str, str]]) -> None:
    ordered = sorted(rows, key=lambda r: SEVERITY_ORDER.get(r.get("flag", ""), 3))
    write_table(ws, sheet.columns, sheet.decisions, sheet.edits, ordered)


def build_pack(catalogue: Catalogue, audit: Path, out: Path) -> dict[str, int]:
    """Write the workbook; return the number of rows per sheet."""
    workbook = Workbook()
    readme = workbook.create_sheet("Read me")
    for default in list(workbook.worksheets[:-1]):
        workbook.remove(default)
    for line in README_LINES:
        readme.append([line])
    readme.column_dimensions["A"].width = 120
    counts = {}
    for sheet in SHEETS:
        rows = ROW_BUILDERS[sheet.title](catalogue, audit)
        _write_sheet(workbook.create_sheet(sheet.title), sheet, rows)
        counts[sheet.title] = len(rows)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return counts


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)
    build = sub.add_parser("build")
    build.add_argument("--version", type=int, required=True)
    build.add_argument("--dir", type=Path, default=DEFAULT_DIR)
    build.add_argument("--audit", type=Path)
    build.add_argument("--out", type=Path)
    apply = sub.add_parser("apply")
    apply.add_argument("workbook", type=Path)
    apply.add_argument("--dir", type=Path, default=DEFAULT_DIR)
    apply.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    if args.command == "apply":
        from scripts.catalogue_review_apply import run_apply

        return run_apply(args.workbook, args.dir, dry_run=args.dry_run)
    pack_dir = REVIEW_DIR / f"catalogue-v{args.version}"
    out = args.out or pack_dir / f"catalogue-review-v{args.version}.xlsx"
    counts = build_pack(read_catalogue(args.dir), args.audit or pack_dir / "audit", out)
    print(f"Wrote {out}: " + ", ".join(f"{k} {v}" for k, v in counts.items()) + ".")
    return 0


if __name__ == "__main__":
    sys.exit(main())
