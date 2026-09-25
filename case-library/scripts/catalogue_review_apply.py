"""Apply Atul's decisions from a returned catalogue review pack to catalogue/*.csv (PLAN L0.5).

Refuses, without writing anything, while any row is undecided or a decision is
not valid for its sheet. Questions are listed for Claude to act on by hand.
"""

import csv
import sys
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from scripts.catalogue import FILES, Catalogue, read_catalogue
from scripts.catalogue_checks import check
from scripts.catalogue_review import SHEETS, range_refs
from scripts.workbook import cell_text

Rows = list[dict[str, str]]


class ReviewError(ValueError):
    """The returned pack cannot be applied. The message lists every problem."""


@dataclass
class Outcome:
    """What applying the pack changed, per file, plus the answers to the questions."""

    changed: dict[str, int] = field(default_factory=dict)
    answers: list[tuple[str, str, str]] = field(default_factory=list)


def read_decisions(workbook_path: Path) -> dict[str, list[dict[str, str]]]:
    """Every sheet's rows as {header: text}; blank rows dropped."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheets: dict[str, list[dict[str, str]]] = {}
    for sheet in SHEETS:
        if sheet.title not in workbook.sheetnames:
            raise ReviewError(f"The workbook has no {sheet.title!r} sheet.")
        values = list(workbook[sheet.title].iter_rows(values_only=True))
        header = [cell_text(v) for v in values[0]]
        rows = [dict(zip(header, map(cell_text, row), strict=False)) for row in values[1:]]
        sheets[sheet.title] = [r for r in rows if r.get(sheet.key)]
    return sheets


def validate(sheets: Mapping[str, list[dict[str, str]]]) -> list[str]:
    problems = []
    for sheet in SHEETS:
        for row in sheets[sheet.title]:
            decision = row.get("decision", "")
            where = f"{sheet.title} {row[sheet.key]}"
            if not decision:
                problems.append(f"{where}: no decision")
            elif sheet.decisions and decision not in sheet.decisions:
                problems.append(f"{where}: decision {decision!r} is not one of {sheet.decisions}")
            elif decision == "edit" and not any(row.get(e) for e in sheet.edits):
                problems.append(f"{where}: 'edit' but no New column is filled")
            elif decision == "suggestion" and not row.get("suggestion"):
                problems.append(f"{where}: 'suggestion' but there is no suggested text")
    return problems


def _rows(catalogue: Catalogue, file: str) -> Rows:
    return [dict(r) for r in catalogue.rows(file)]


def _set(row: dict[str, str], column: str, value: str) -> bool:
    if value and row[column] != value:
        row[column] = value
        return True
    return False


def _apply_ranges(
    catalogue: Catalogue, decided: list[dict[str, str]], files: dict[str, Rows]
) -> None:
    by_ref = {d["ref"]: d for d in decided if d["decision"] == "edit"}
    components = {r["id"]: r for r in files["components.csv"]}
    ranges = files["component_ranges.csv"]
    for index, (ref, _) in enumerate(range_refs(catalogue)):
        edit = by_ref.get(ref)
        if edit is None:
            continue
        for column in ("low", "high", "display", "text", "source"):
            _set(ranges[index], column, edit.get(f"new_{column}", ""))
        component = components[ranges[index]["component_id"]]
        for column in ("unit_si", "decimals", "conv_factor"):
            _set(component, column, edit.get(f"new_{column}", ""))


def _apply_texts(decided: list[dict[str, str]], files: dict[str, Rows]) -> None:
    templates = {r["item_id"]: r for r in files["normal_templates.csv"]}
    components = {r["id"]: r for r in files["components.csv"]}
    for d in decided:
        text = {"edit": d.get("new_text", ""), "suggestion": d.get("suggestion", "")}.get(
            d["decision"], ""
        )
        if d["item_id"] in templates:
            row = templates[d["item_id"]]
            _set(row, "template", text)
            row["review_status"] = "approved"
        elif d["item_id"] in components:
            _set(components[d["item_id"]], "normal_text", text)


def _apply_prices(decided: list[dict[str, str]], files: dict[str, Rows]) -> None:
    tests = {r["id"]: r for r in files["tests.csv"]}
    for d in decided:
        row = tests[d["test_id"]]
        if d["decision"] == "approve":
            _set(row, "price_inr", d["proposed_inr"])
            _set(row, "price_source", d["proposed_source"])
            continue
        _set(row, "price_inr", d.get("new_price_inr", ""))
        source = d.get("new_price_source", "")
        if d.get("new_price_inr") and not source:
            source = "estimate" if row["price_source"] == "estimate" else "reviewer"
        _set(row, "price_source", source)
        _set(row, "tat_minutes", d.get("new_tat_minutes", ""))


def _apply_codes(decided: list[dict[str, str]], files: dict[str, Rows]) -> None:
    diagnoses = {r["id"]: r for r in files["diagnoses.csv"]}
    for d in decided:
        row = diagnoses[d["dx_id"]]
        edit = d["decision"] == "edit"
        _set(row, "icd10", d.get("new_icd10", "") if edit else d["proposed_icd10"])
        new_icd11 = d.get("new_icd11", "") if edit else d["proposed_icd11"]
        if new_icd11 or not edit:
            row["icd11"] = new_icd11 or row["icd11"]
        row["codes_verified"] = "false" if d["decision"] == "unverified" else "true"


def apply_decisions(
    catalogue: Catalogue, sheets: Mapping[str, list[dict[str, str]]]
) -> tuple[dict[str, Rows], Outcome]:
    """The new file contents and what changed; raises ReviewError if the pack is incomplete."""
    problems = validate(sheets)
    if problems:
        raise ReviewError("\n".join(problems))
    files = {name: _rows(catalogue, name) for name in FILES}
    _apply_ranges(catalogue, sheets["Ranges"], files)
    _apply_texts(sheets["Normal texts"], files)
    _apply_prices(sheets["Prices and turnaround"], files)
    _apply_codes(sheets["Diagnosis codes"], files)
    outcome = Outcome()
    for name, rows in files.items():
        changed = sum(
            1 for old, new in zip(catalogue.rows(name), rows, strict=True) if dict(old) != new
        )
        if changed:
            outcome.changed[name] = changed
    outcome.answers = [
        (q["question_id"], q["question"], q["decision"]) for q in sheets["Questions"]
    ]
    return files, outcome


def write_files(directory: Path, files: Mapping[str, Rows]) -> None:
    for name, rows in files.items():
        with (directory / name).open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(FILES[name]), lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)


def run_apply(workbook: Path, directory: Path, *, dry_run: bool) -> int:
    catalogue = read_catalogue(directory)
    try:
        files, outcome = apply_decisions(catalogue, read_decisions(workbook))
    except ReviewError as err:
        print(f"Not applied:\n{err}", file=sys.stderr)
        return 1
    after = check(Catalogue({n: tuple(r) for n, r in files.items()}))
    if after:
        print(
            "Not applied: the decisions leave the catalogue broken:\n" + "\n".join(after),
            file=sys.stderr,
        )
        return 1
    if not dry_run:
        write_files(directory, files)
    verb = "Would change" if dry_run else "Changed"
    print(
        f"{verb}: "
        + (", ".join(f"{k} {v} rows" for k, v in outcome.changed.items()) or "nothing")
        + "."
    )
    for question_id, question, answer in outcome.answers:
        print(f"Question {question_id}: {question}\n  Answer: {answer}")
    return 0
