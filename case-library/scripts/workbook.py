"""Shared styling for the Excel review packs Atul fills in (catalogue and case packs).

Each review sheet shows its columns, then a yellow Decision column with a dropdown,
the New columns an edit fills, and a free Note column.
"""

from collections.abc import Callable, Mapping, Sequence
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

FILLS = {
    "error": PatternFill("solid", fgColor="F8CBAD"),
    "check": PatternFill("solid", fgColor="FFE699"),
    "decision": PatternFill("solid", fgColor="FFF2CC"),
    "header": PatternFill("solid", fgColor="D9E1F2"),
}
WIDE_COLUMNS = frozenset(
    {
        "text",
        "suggestion",
        "issue",
        "question",
        "options",
        "cghs_names",
        "match_note",
        "audit_note",
        "recommendation",
        "new_text",
        "note",
        "rationale",
        "release_text",
        "lay_text",
        "status_line",
        "condition",
        "correction",
    }
)
WIDE, NARROW = 60, 16

Row = Mapping[str, str]


def _flag_of(row: Row) -> str:
    return row.get("flag", "")


def write_table(
    ws: Worksheet,
    columns: Sequence[str],
    decisions: Sequence[str],
    edits: Sequence[str],
    rows: Sequence[Row],
    *,
    fill_of: Callable[[Row], str] = _flag_of,
) -> None:
    """Write rows in the order given; the first cell takes the fill named by fill_of(row)."""
    header = [*columns, "decision", *edits, "note"]
    ws.append(header)
    for row in rows:
        ws.append([row.get(col, "") for col in columns] + [""] * (len(header) - len(columns)))
        fill = FILLS.get(fill_of(row))
        if fill is not None:
            ws.cell(ws.max_row, 1).fill = fill
    first_decision = len(columns) + 1
    _style_header(ws, header, first_decision)
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row_cells[first_decision - 1].fill = FILLS["decision"]
    ws.freeze_panes = "B2"
    if decisions and ws.max_row > 1:
        validation = DataValidation(
            type="list", formula1=f'"{",".join(decisions)}"', allow_blank=True
        )
        ws.add_data_validation(validation)
        letter = get_column_letter(first_decision)
        validation.add(f"{letter}2:{letter}{ws.max_row}")


def _style_header(ws: Worksheet, header: Sequence[str], first_decision: int) -> None:
    for col_index, name in enumerate(header, start=1):
        cell = ws.cell(1, col_index)
        cell.font = Font(bold=True)
        cell.fill = FILLS["decision"] if col_index >= first_decision else FILLS["header"]
        width = WIDE if name in WIDE_COLUMNS else NARROW
        ws.column_dimensions[get_column_letter(col_index)].width = width


def cell_text(value: Any) -> str:
    """A workbook cell as stripped text; whole-number floats lose their '.0'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()
