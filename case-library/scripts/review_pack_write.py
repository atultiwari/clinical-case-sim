"""Write a case review pack workbook from PackData (SPEC §7.2).

One README sheet, then one sheet per review area. Rows needing most attention come
first (judgement calls have their own sheet; then high priority). Atul fills the
yellow Decision column from its dropdown, the New columns for an edit, and a note.
"""

from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from scripts.review_pack_data import CaseSummary, PackData, Row
from scripts.review_pack_sql import (
    ARTICLE_FACT_COLUMNS,
    NORMAL_LIST_COLUMNS,
    PATIENT_WORDS_COLUMNS,
    REPORT_COLUMNS,
)
from scripts.workbook import write_table

APPROVE = ("Approve", "Edit", "Reject")
CONFIRM = ("OK", "Correction")
ORIGINS = ("article", "derived", "affected", "normal", "rule", "reviewer")
PRIORITY_ORDER = {"high": 0, "medium": 1, "low": 2}
AFFECTED_COLUMNS = (
    "key",
    "case",
    "target",
    "item",
    "tier",
    "day",
    "value",
    "ref_range",
    "release_text",
    "rationale",
    "confidence",
    "priority",
)


def _all(_: Row) -> bool:
    return True


def _is_judgement(row: Row) -> bool:
    return row.get("judgement_call") == "yes"


def _not_judgement(row: Row) -> bool:
    return not _is_judgement(row)


@dataclass(frozen=True)
class PackSheet:
    """One review sheet: where its rows come from, what it shows and how it is decided."""

    title: str
    source: str  # a PackData field
    columns: tuple[str, ...]
    decisions: tuple[str, ...]
    edits: tuple[str, ...] = ()
    keep: Callable[[Row], bool] = _all
    fill: str = ""  # a fill for every row (the leak scan), otherwise by priority

    @property
    def needs_edit(self) -> tuple[str, ...]:
        """The decisions that need one of the edit columns filled."""
        return tuple(d for d in self.decisions if d in {"Edit", "Correction"})


PACK_SHEETS: tuple[PackSheet, ...] = (
    PackSheet(
        "Judgement calls",
        "ledger",
        AFFECTED_COLUMNS,
        APPROVE,
        ("new_value", "new_text"),
        _is_judgement,
    ),
    PackSheet(
        "Affected", "ledger", AFFECTED_COLUMNS, APPROVE, ("new_value", "new_text"), _not_judgement
    ),
    PackSheet("Normal list", "normal_list", NORMAL_LIST_COLUMNS, ("OK", "Should be affected")),
    PackSheet("Article facts", "article_facts", ARTICLE_FACT_COLUMNS, CONFIRM, ("correction",)),
    PackSheet("Patient's words", "patient_words", PATIENT_WORDS_COLUMNS, CONFIRM, ("correction",)),
    PackSheet("Reports and consults", "reports_consults", REPORT_COLUMNS, APPROVE, ("new_text",)),
    PackSheet(
        "Ground truth",
        "ground_truth",
        ("key", "case", "part", "text", "condition"),
        CONFIRM,
        ("correction",),
    ),
    PackSheet(
        "Leak scan",
        "leaks",
        ("key", "case", "location", "row_id", "term"),
        ("Resolved",),
        fill="error",
    ),
)
SHEET_BY_TITLE: Mapping[str, PackSheet] = {s.title: s for s in PACK_SHEETS}

HOW_TO = (
    "How to fill it in",
    "Every row needs a Decision (the yellow column, from its dropdown). Fill-down is fine for a "
    "block of rows you accept. The key column links each row to the Case Vault; leave it as it is.",
    "Judgement calls and Affected: Approve, Reject, or Edit with new_value (and new_text for "
    "the wording shown to the player). Amber rows are judgement calls or high priority.",
    "Normal list: the items the normal generator or a rule answers (values are not shown). OK, "
    "or 'Should be affected' if the diagnosis, a comorbidity or a medicine would change it.",
    "Article facts: check each value against its source locator in the article. OK, or "
    "Correction with the right value in the correction column.",
    "Patient's words: the lay text must carry exactly the clinical fact, with no added symptoms "
    "and no lost negatives. OK, or Correction with the new lay text.",
    "Reports and consults: Approve, Reject, or Edit with new_text. A provisional report must "
    "say it is not final.",
    "Ground truth: OK, or Correction with the new wording or condition.",
    "Leak scan: player-visible text that names the diagnosis. Mark Resolved once you have "
    "decided what to do; say what in the note.",
    "Return the file as it is; Claude reads it with `python -m scripts.review_pack read` and "
    "shows you a summary before writing anything to the Case Vault.",
)


def _order(row: Row) -> tuple[int, int]:
    return (0 if _is_judgement(row) else 1, PRIORITY_ORDER.get(row.get("priority", ""), 3))


def _fill_for(sheet: PackSheet) -> Callable[[Row], str]:
    def fill(row: Row) -> str:
        if sheet.fill:
            return sheet.fill
        return "check" if _order(row) < (1, 1) else ""

    return fill


def sheet_rows(data: PackData, sheet: PackSheet) -> list[Row]:
    rows: Sequence[Row] = getattr(data, sheet.source)
    return sorted((r for r in rows if sheet.keep(r)), key=_order)


def _flag(value: bool | None) -> str:
    return "n/a" if value is None else ("yes" if value else "no")


def _case_line(case: CaseSummary) -> list[str | int]:
    head: list[str | int] = [
        case.case_version_id,
        case.status,
        case.source,
        case.licence,
        _flag(case.production_ok),
        _flag(case.public_release_ok),
    ]
    return head + [case.origin_counts.get(origin, 0) for origin in ORIGINS]


def _write_readme(ws: Worksheet, data: PackData, counts: Mapping[str, int]) -> None:
    ws.append(["Case review pack (Case Library, SPEC §7.2). Educational and research use only."])
    ws.append(["It shows diagnoses: keep it private and out of Git."])
    ws.append(["Batch", data.batch])
    ws.append([])
    ws.append(
        ["case", "status", "source", "licence", "production_ok", "public_release_ok", *ORIGINS]
    )
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for case in data.cases:
        ws.append(_case_line(case))
    ws.append([])
    ws.append(["Sheet", "rows"])
    for title, count in counts.items():
        ws.append([title, count])
    ws.append([])
    for line in HOW_TO:
        ws.append([line])
    ws.column_dimensions["A"].width = 30
    for letter in "BCDEF":
        ws.column_dimensions[letter].width = 18


def write_pack(data: PackData, out: Path) -> dict[str, int]:
    """Write the workbook; return the number of rows per review sheet."""
    workbook = Workbook()
    readme = workbook.worksheets[0]
    readme.title = "README"
    counts: dict[str, int] = {}
    for sheet in PACK_SHEETS:
        rows = sheet_rows(data, sheet)
        ws = workbook.create_sheet(sheet.title)
        write_table(ws, sheet.columns, sheet.decisions, sheet.edits, rows, fill_of=_fill_for(sheet))
        ws.column_dimensions["A"].outline_level = 1  # the key column can be folded away
        counts[sheet.title] = len(rows)
    _write_readme(readme, data, counts)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return counts
