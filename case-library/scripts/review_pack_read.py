"""Read a returned case review pack (SPEC §7.3). It never writes to the Case Vault.

Refuses, listing every problem, while any row is undecided, a decision is not one
of its sheet's choices, or an Edit or Correction leaves its New columns empty.
Otherwise it summarises the decisions per case and writes them, one per row, to
<pack>.decisions.json beside the workbook. Claude applies them through the MCP
after Atul's go-ahead.
"""

import json
import sys
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

from scripts.review_pack_data import ReviewPackError
from scripts.review_pack_write import PACK_SHEETS, PackSheet
from scripts.workbook import cell_text

SheetRows = Mapping[str, list[dict[str, str]]]
SUMMARY_LABELS = (
    ("approve", "approved"),
    ("edit", "edited"),
    ("reject", "rejected"),
    ("correction", "corrections"),
    ("ok", "OK"),
    ("resolved", "leaks resolved"),
)


@dataclass(frozen=True)
class Decision:
    key: str
    sheet: str
    table: str
    id: str
    case_version_id: str
    decision: str  # approve, edit, reject, ok, correction, should_be_affected, resolved
    edited: dict[str, str] | None
    note: str


def read_pack(path: Path) -> dict[str, list[dict[str, str]]]:
    """Every review sheet's rows as {header: text}; rows without a key are dropped."""
    if not path.exists():
        raise ReviewPackError(f"{path} not found.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list[dict[str, str]]] = {}
    for sheet in PACK_SHEETS:
        if sheet.title not in workbook.sheetnames:
            raise ReviewPackError(f"The workbook has no {sheet.title!r} sheet.")
        values = list(workbook[sheet.title].iter_rows(values_only=True))
        header = [cell_text(v) for v in values[0]] if values else []
        rows = [dict(zip(header, map(cell_text, row), strict=False)) for row in values[1:]]
        sheets[sheet.title] = [r for r in rows if r.get("key")]
    workbook.close()
    return sheets


def _choice(sheet: PackSheet, decision: str) -> str | None:
    """The sheet's own spelling of a decision, in any letter case; None if not one of them."""
    return next((d for d in sheet.decisions if d.lower() == decision.lower()), None)


def _problem(sheet: PackSheet, row: Mapping[str, str]) -> str | None:
    where = f"{sheet.title} {row['key']}"
    decision = row.get("decision", "")
    choice = _choice(sheet, decision)
    if not decision:
        return f"{where}: no decision"
    if choice is None:
        return f"{where}: decision {decision!r} is not one of {', '.join(sheet.decisions)}"
    if choice in sheet.needs_edit and not any(row.get(e) for e in sheet.edits):
        return f"{where}: {choice} but no {' or '.join(sheet.edits)}"
    return None


def validate(sheets: SheetRows) -> list[str]:
    problems = []
    for sheet in PACK_SHEETS:
        for row in sheets[sheet.title]:
            problem = _problem(sheet, row)
            if problem is not None:
                problems.append(problem)
    return problems


def _decision(sheet: PackSheet, row: Mapping[str, str]) -> Decision:
    table, _, target_id = row["key"].partition(":")
    choice = _choice(sheet, row["decision"]) or ""
    edited = {e: row[e] for e in sheet.edits if row.get(e)} if choice in sheet.needs_edit else {}
    return Decision(
        key=row["key"],
        sheet=sheet.title,
        table=table,
        id=target_id,
        case_version_id=row.get("case", ""),
        decision=choice.lower().replace(" ", "_"),
        edited=edited or None,
        note=row.get("note", ""),
    )


def collect_decisions(sheets: SheetRows) -> list[Decision]:
    """The decisions in sheet order; raises ReviewPackError listing every problem."""
    problems = validate(sheets)
    if problems:
        raise ReviewPackError("\n".join(problems))
    return [_decision(sheet, row) for sheet in PACK_SHEETS for row in sheets[sheet.title]]


def summarise(decisions: Sequence[Decision]) -> list[str]:
    """One line per case: approved, edited, rejected, corrections (incl. 'should be affected')."""
    counts: dict[str, Counter[str]] = {}
    for d in decisions:
        kind = "correction" if d.decision == "should_be_affected" else d.decision
        counts.setdefault(d.case_version_id, Counter())[kind] += 1
    return [
        f"{case}: " + ", ".join(f"{label} {counter[kind]}" for kind, label in SUMMARY_LABELS)
        for case, counter in sorted(counts.items())
    ]


def decisions_path(pack: Path) -> Path:
    return pack.with_name(f"{pack.stem}.decisions.json")


def run_read(pack: Path) -> int:
    try:
        decisions = collect_decisions(read_pack(pack))
    except ReviewPackError as err:
        print(f"Not read:\n{err}", file=sys.stderr)
        return 1
    out = decisions_path(pack)
    body = json.dumps([asdict(d) for d in decisions], ensure_ascii=False, indent=2)
    out.write_text(body + "\n", encoding="utf-8")
    for line in summarise(decisions):
        print(line)
    print(f"Wrote {len(decisions)} decisions to {out}. Nothing has been written to the Case Vault.")
    return 0
